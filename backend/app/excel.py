"""Excel report generation (openpyxl). Group A and Group B are always separate workbooks."""
from __future__ import annotations

import io
from collections import OrderedDict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import statuses as S
from .config import GROUP_A, GROUP_B

MONEY_FMT = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="1F3A68")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D7E2")


def _d(v) -> Decimal | None:
    return None if v is None else Decimal(str(v))


def _sheet(ws, headers: list[str], rows: list[list], money_cols: set[int] = frozenset()):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    widths = [len(h) for h in headers]
    for row in ws.iter_rows(min_row=1):
        for i, c in enumerate(row):
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if i in money_cols and c.row > 1 and isinstance(c.value, (int, float, Decimal)):
                c.number_format = MONEY_FMT
            widths[i] = max(widths[i], len(str(c.value or "")) + 2)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 10), 60)
    ws.freeze_panes = "A2"


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_group_a(files) -> bytes:
    files = [f for f in files if f.group == GROUP_A]
    ok = [f for f in files if f.status == S.PROCESSED]
    wb = Workbook()

    ws = wb.active
    ws.title = "Allocation Summary"
    rows = []
    for f in ok:
        for a in f.allocations:
            rows.append([f.file_name, f.contract_number, a.get("subscription") or f.subscription_number, f.bill_period, f.bill_month,
                         a["company"], a["dimension"], Decimal(a["amount"])])
    _sheet(ws, ["Source File Name", "Contract Number", "Subscription Number", "Bill Period", "Bill Month",
                "Company", "Dimension", "Allocated Amount"], rows, {7})

    ws = wb.create_sheet("Calculation Summary")
    rows = []
    multi = len(ok) > 1
    for f in ok:
        if multi:
            rows.append([f"Source: {f.file_name} ({f.bill_month})", None])
        rows += [["Total Charges for Bill Period", _d(f.total_charges)], ["VAT", _d(f.vat)],
                 ["Total Without VAT", _d(f.total_without_vat)],
                 ["Allocation Per Company", _d(f.allocation_per_company).quantize(Decimal("0.0001"))]
                 if f.allocation_per_company is not None else
                 ["Allocation Per Company", "Per subscription (itemised on bill)"]]
        adj = _d(f.rounding_adjustment) or Decimal(0)
        if adj:
            rows.append([f"Rounding adjustment applied to {f.allocations[-1]['company']}", adj])
    if multi:
        rows.append(["Batch Total - Total Charges for Bill Period", sum(_d(f.total_charges) for f in ok)])
        rows.append(["Batch Total - VAT", sum(_d(f.vat) for f in ok)])
        rows.append(["Batch Total - Total Without VAT", sum(_d(f.total_without_vat) for f in ok)])
    _sheet(ws, ["Description", "Amount"], rows, {1})
    ws.append([])
    ws.append(["Note: allocations are rounded to 2 decimals (half-up); any remainder is applied to the "
               "last company so allocations sum exactly to Total Without VAT."])

    ws = wb.create_sheet("Validation")
    # Only two statuses are permitted in this sheet; any non-success outcome is a File Read Failure.
    rows = [[f.file_name, f.contract_number,
             S.PROCESSED if f.status == S.PROCESSED else S.READ_FAILURE] for f in files]
    _sheet(ws, ["Source File Name", "Contract Number", "Status"], rows)
    return _to_bytes(wb)


def build_group_b(files) -> bytes:
    files = [f for f in files if f.group == GROUP_B]
    ok = [f for f in files if f.status == S.PROCESSED]
    wb = Workbook()

    ws = wb.active
    ws.title = "Bill Summary"
    rows = [[f.file_name, f.contract_number, f.subscription_number, f.bill_period, f.bill_month,
             f.related_company, f.cost_center, _d(f.total_charges), _d(f.vat), _d(f.total_without_vat), f.status]
            for f in ok]
    _sheet(ws, ["Source File Name", "Contract Number", "Subscription Number", "Bill Period", "Bill Month",
                "Related Company", "Cost Center", "Total Charges for Bill Period", "VAT",
                "Total Charges without VAT", "Status"], rows, {7, 8, 9})

    ws = wb.create_sheet("Cost Center Summary")
    totals: "OrderedDict[str, list]" = OrderedDict()
    for f in ok:
        entry = totals.setdefault(f.cost_center, [set(), Decimal(0)])
        entry[0].add(f.related_company)
        entry[1] += _d(f.total_without_vat)
    rows = [[cc, " / ".join(sorted(companies)), amt] for cc, (companies, amt) in totals.items()]
    _sheet(ws, ["Cost Center", "Company", "Total Amount"], rows, {2})

    ws = wb.create_sheet("Failed Records")
    rows = []
    for f in files:
        if f.status == S.PROCESSED:
            continue
        status = f.status if not f.errors else f"{f.status}: {'; '.join(f.errors)}"
        rows.append([f.file_name, f.contract_number, status])
    _sheet(ws, ["Source File Name", "Contract Number", "Status"], rows)
    return _to_bytes(wb)


