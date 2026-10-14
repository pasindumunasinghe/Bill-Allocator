"""Usage file reading, username matching (Master data -> map.xlsx) and cost allocation."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl

CENT = Decimal("0.01")
EXCLUDED = {"total", "others", "other", "grandtotal"}
UNIT_TO_GB = {"tb": Decimal(1024), "gb": Decimal(1), "mb": Decimal(1) / 1024,
              "kb": Decimal(1) / 1024 ** 2, "bytes": Decimal(1) / 1024 ** 3, "b": Decimal(1) / 1024 ** 3}


class UsageError(Exception):
    pass


# ---------- username helpers ----------

def raw_user(value) -> str:
    """'OREL\\kamal' -> 'kamal', 'Kamal.P@orel.com' -> 'kamal.p' (trimmed, lowercase)."""
    s = str(value or "").strip().lower()
    if "\\" in s:
        s = s.rsplit("\\", 1)[1]
    if "@" in s:
        s = s.split("@", 1)[0]
    return s.strip()


def norm(value) -> str:
    """Matching key: raw_user with '.', '_', '-' and spaces removed."""
    return re.sub(r"[._\-\s]", "", raw_user(value))


def is_excluded(user: str) -> bool:
    n = norm(user)
    return not n or n in EXCLUDED or n.startswith("total")


# ---------- generic sheet reading ----------

def _header_row(rows: list[tuple], must: list[list[str]], scan: int = 20):
    """Find the first row containing a header matching each group of keywords. Returns (index, {group: col})."""
    for i, row in enumerate(rows[:scan]):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        found = {}
        for g, keys in enumerate(must):
            for j, c in enumerate(cells):
                if c and j not in found.values() and any(k(c) if callable(k) else k in c for k in keys):
                    found[g] = j
                    break
        if len(found) == len(must):
            return i, found
    return None, None


def _number(v) -> tuple[Decimal | None, str | None]:
    """Numeric cell or text like '1.5 GB'. Returns (value, unit-in-text)."""
    if v is None or v == "":
        return None, None
    if isinstance(v, (int, float)):
        return Decimal(str(v)), None
    m = re.match(r"^\s*(-?[\d,]*\.?\d+)\s*([A-Za-z]+)?\s*$", str(v))
    if not m:
        return None, None
    try:
        return Decimal(m.group(1).replace(",", "")), (m.group(2) or "").lower() or None
    except InvalidOperation:
        return None, None


def _unit(header: str) -> str | None:
    h = header.lower()
    for u in ("tb", "gb", "mb", "kb", "bytes"):
        if re.search(rf"(?<![a-z]){u}(?![a-z])", h):
            return u
    return None


@dataclass
class SheetUsage:
    name: str
    kind: str  # "source_user" (grouped sum) or "username" (first occurrence)
    user_col: str
    usage_col: str
    unit: str
    usage: dict[str, Decimal] = field(default_factory=dict)  # raw username -> GB
    excluded_rows: int = 0
    duplicates_removed: int = 0
    notes: list[str] = field(default_factory=list)


USAGE_KEYS = ["usage", "total", "bytes", "volume", "data", "gb", "mb"]


def read_usage_workbook(path: str) -> list[SheetUsage]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise UsageError(f"Usage file could not be read: {exc}") from exc
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        idx, cols = _header_row(rows, [["source user", "sourceuser"], USAGE_KEYS])
        kind = "source_user"
        if idx is None:
            idx, cols = _header_row(rows, [[lambda c: c.replace("_", " ").replace(" ", "") in ("username", "user", "user name")
                                            or c.startswith("username")], USAGE_KEYS])
            kind = "username"
        if idx is None:
            continue  # not a usage sheet
        header = rows[idx]
        ucol, vcol = cols[0], cols[1]
        sh = SheetUsage(ws.title, kind, str(header[ucol]), str(header[vcol]), _unit(str(header[vcol])) or "")
        seen = set()
        for r in rows[idx + 1:]:
            if ucol >= len(r) or r[ucol] in (None, ""):
                continue
            user = raw_user(r[ucol])
            if is_excluded(user):
                sh.excluded_rows += 1
                continue
            val, text_unit = _number(r[vcol] if vcol < len(r) else None)
            if val is None:
                continue
            unit = text_unit or sh.unit
            if unit and unit not in UNIT_TO_GB:
                raise UsageError(f"Sheet '{ws.title}': unknown usage unit '{unit}'")
            gb = val * UNIT_TO_GB[unit] if unit else val
            if kind == "username":
                key = norm(user)
                if key in seen:  # keep only the first occurrence
                    sh.duplicates_removed += 1
                    continue
                seen.add(key)
                sh.usage[user] = gb
            else:
                sh.usage[user] = sh.usage.get(user, Decimal(0)) + gb
        if not sh.unit:
            sh.notes.append(f"No unit in column '{sh.usage_col}'; values used as GB")
        sheets.append(sh)
    wb.close()
    if not sheets:
        raise UsageError("No usage sheet found (need a 'Source User' or 'username' column and a usage column)")
    return sheets


# ---------- master / map ----------

@dataclass
class Person:
    emp_no: str
    emp_name: str
    company: str
    dimension: str
    key: str  # normalized username / email prefix
    source: str


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_master(path: str) -> list[Person]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    people = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        idx, cols = _header_row(rows, [["emp no", "empno", "emp_no", "employee no", "employee number", "emp id", "epf"],
                                       ["company"], ["dimension", "cost center", "cost centre"]])
        if idx is None:
            continue
        head = [str(c or "").strip().lower() for c in rows[idx]]
        email = next((j for j, c in enumerate(head) if "mail" in c), None)
        uname = next((j for j, c in enumerate(head) if "user" in c), None)
        name = next((j for j, c in enumerate(head) if "name" in c and "user" not in c), None)
        for r in rows[idx + 1:]:
            get = lambda j: _s(r[j]) if j is not None and j < len(r) else ""  # noqa: E731
            if not get(cols[0]):
                continue
            keys = {norm(get(email))} if get(email) else set()
            if get(uname):
                keys.add(norm(get(uname)))
            for k in keys - {""}:
                people.append(Person(get(cols[0]), get(name), get(cols[1]), get(cols[2]), k, "Master data"))
    wb.close()
    if not people:
        raise UsageError("Master data: no rows with Emp No / Company / Dimension columns found")
    return people


def read_map(path: str) -> list[Person]:
    """map.xlsx: Username, Emp No, Emp Name, Company (3rd column), Dimension."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    people = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        idx, cols = _header_row(rows, [["user"], ["emp no", "empno", "emp_no", "employee no", "emp id", "epf"]])
        if idx is None:
            continue
        head = [str(c or "").strip().lower() for c in rows[idx]]
        name = next((j for j, c in enumerate(head) if "name" in c and "user" not in c), None)
        comp = next((j for j, c in enumerate(head) if "company" in c), 2)
        dim = next((j for j, c in enumerate(head) if "dimension" in c or "cost cent" in c), None)
        for r in rows[idx + 1:]:
            get = lambda j: _s(r[j]) if j is not None and j < len(r) else ""  # noqa: E731
            if get(cols[0]):
                people.append(Person(get(cols[1]), get(name), get(comp), get(dim), norm(get(cols[0])), "map.xlsx"))
    wb.close()
    return people


def match(key: str, people: list[Person]) -> tuple[Person | None, str | None]:
    """Exact normalized match first; otherwise a partial match only if it is unique (never guess)."""
    exact = [p for p in people if p.key == key]
    if exact:
        return exact[0], "exact"
    if len(key) >= 4:
        partial = {(p.emp_no, p.key): p for p in people if len(p.key) >= 4 and (key in p.key or p.key in key)}
        if len(partial) == 1:
            return next(iter(partial.values())), "partial"
    return None, None


# ---------- allocation ----------

@dataclass
class UserLine:
    raw: str
    key: str
    per_sheet: dict[str, Decimal]
    person: Person | None = None
    matched_by: str = ""
    total: Decimal = Decimal(0)
    share: Decimal = Decimal(0)  # redistributed All Company usage
    final: Decimal = Decimal(0)
    cost: Decimal = Decimal(0)

    @property
    def all_company(self) -> bool:
        return bool(self.person) and norm(self.person.emp_no) == "allcompany"


@dataclass
class Allocation:
    sheets: list[SheetUsage]
    users: list[UserLine]  # mapped users, including All Company
    unmapped: list[UserLine]
    pool: Decimal
    recipients: int
    grand_total: Decimal
    unit_cost: Decimal
    rounding_adjustment: Decimal
    allocated: Decimal


def allocate(bill_value: Decimal, sheets: list[SheetUsage], master: list[Person], mapping: list[Person]) -> Allocation:
    lines: dict[str, UserLine] = {}
    for sh in sheets:
        for user, gb in sh.usage.items():
            key = norm(user)
            ln = lines.setdefault(key, UserLine(user, key, {}))
            ln.per_sheet[sh.name] = ln.per_sheet.get(sh.name, Decimal(0)) + gb
    mapped, unmapped = [], []
    for ln in lines.values():
        ln.total = sum(ln.per_sheet.values(), Decimal(0))
        p, how = match(ln.key, master)
        if p:
            ln.person, ln.matched_by = p, f"Master data ({how})"
        else:
            p, how = match(ln.key, mapping)
            if p:
                ln.person, ln.matched_by = p, f"map.xlsx ({how})"
        (mapped if ln.person else unmapped).append(ln)

    pool = sum((u.total for u in mapped if u.all_company), Decimal(0))
    recipients = [u for u in mapped if not u.all_company]
    if not recipients:
        raise UsageError("No mapped users to allocate the bill to")
    share = pool / len(recipients)
    for u in mapped:
        u.share = Decimal(0) if u.all_company else share
        u.final = Decimal(0) if u.all_company else u.total + share
    grand = sum((u.final for u in recipients), Decimal(0))
    if grand <= 0:
        raise UsageError("Total usage is zero; cannot allocate")
    unit_cost = bill_value / grand
    for u in recipients:
        u.cost = (unit_cost * u.final).quantize(CENT, ROUND_HALF_UP)
    adj = bill_value.quantize(CENT, ROUND_HALF_UP) - sum(u.cost for u in recipients)
    max(recipients, key=lambda u: u.final).cost += adj  # rounding remainder -> largest user
    allocated = sum(u.cost for u in recipients)
    if allocated != bill_value.quantize(CENT, ROUND_HALF_UP):
        raise UsageError("Allocated cost does not equal bill value")
    mapped.sort(key=lambda u: (u.person.company, u.raw))
    unmapped.sort(key=lambda u: u.raw)
    return Allocation(sheets, mapped, unmapped, pool, len(recipients), grand, unit_cost, adj, allocated)
