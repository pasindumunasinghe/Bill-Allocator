"""Excel output for usage-based allocation: one workbook per bill."""
from __future__ import annotations

import io
import re
from collections import OrderedDict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .bill import BillValue
from .usage import Allocation

FILL = PatternFill("solid", fgColor="1F3A68")
WHITE = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MONEY, GB = "#,##0.00", "#,##0.0000"


def _table(ws, headers, rows, fmt: dict[int, str] | None = None, total_row=None):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.fill, c.font = FILL, WHITE
    for r in rows:
        ws.append(r)
    if total_row:
        ws.append(total_row)
        for c in ws[ws.max_row]:
            c.font = BOLD
    for row in ws.iter_rows(min_row=2):
        for i, c in enumerate(row):
            if fmt and i in fmt and isinstance(c.value, (int, float, Decimal)):
                c.number_format = fmt[i]
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(h)) + 4))
    ws.freeze_panes = "A2"


def build(bill: BillValue, alloc: Allocation) -> tuple[bytes, str]:
    names = [s.name for s in alloc.sheets]
    wb = Workbook()

    ws = wb.active
    ws.title = "Allocation Summary"
    head = ["Raw Username", "Normalized Username", "Emp No", "Emp Name", "Company", "Dimension", "Matched By"] \
        + [f"{n} Usage (GB)" for n in names] + ["Total Usage (GB)", "All Company Share (GB)", "Final Usage (GB)",
                                                 "Unit Cost", "Allocated Cost"]
    rows = []
    for u in alloc.users:
        p = u.person
        rows.append([u.raw, u.key, p.emp_no, p.emp_name, p.company, p.dimension, u.matched_by]
                    + [u.per_sheet.get(n, Decimal(0)) for n in names]
                    + [u.total, u.share, u.final, alloc.unit_cost if not u.all_company else Decimal(0), u.cost])
    n = len(names)
    fmt = {i: GB for i in range(7, 7 + n + 3)} | {7 + n + 3: "#,##0.000000", 7 + n + 4: MONEY}
    _table(ws, head, rows, fmt, ["TOTAL", "", "", "", "", "", ""] + [""] * n
           + [sum((u.total for u in alloc.users), Decimal(0)), "", alloc.grand_total, "", alloc.allocated])

    ws = wb.create_sheet("Unmapped Users")
    _table(ws, ["Username"] + [f"{n} Usage (GB)" for n in names] + ["Total Usage (GB)", "Reason"],
           [[u.raw] + [u.per_sheet.get(n, Decimal(0)) for n in names] + [u.total, "User not found in Master data"]
            for u in alloc.unmapped], {i: GB for i in range(1, n + 2)})

    ws = wb.create_sheet("Bill Summary")
    rows = [["Bill File", bill.file_name], ["Bill Type", bill.bill_type],
            ["Contract / Account Number", bill.reference], ["Bill Period", bill.bill_period or "Not found in bill"]]
    rows += [[label, amt] for label, amt in bill.components]
    rows += [["Bill Value (allocated)", bill.value], ["VAT (excluded, for reference)", bill.vat],
             ["Grand Total Usage (GB, mapped users after All Company redistribution)", alloc.grand_total],
             ["Unit Cost per GB", alloc.unit_cost.quantize(Decimal("0.000001"))],
             ["Total Allocated Cost", alloc.allocated],
             ["Difference (Bill Value - Allocated)", bill.value - alloc.allocated],
             ["Rounding adjustment (applied to highest-usage user)", alloc.rounding_adjustment],
             ["Mapped users", len(alloc.users)], ["Unmapped users (not allocated)", len(alloc.unmapped)],
             ["Unmapped usage (GB, not allocated)", sum((u.total for u in alloc.unmapped), Decimal(0))]]
    for s in alloc.sheets:
        rows.append([f"Sheet '{s.name}'", f"{s.kind.replace('_', ' ')} column '{s.user_col}', usage '{s.usage_col}'"
                     f" ({s.unit or 'GB assumed'}); {s.excluded_rows} total/others rows excluded; "
                     f"{s.duplicates_removed} duplicates removed" + ("; " + "; ".join(s.notes) if s.notes else "")])
    _table(ws, ["Item", "Value"], rows, {1: MONEY})
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 50

    ws = wb.create_sheet("All Company Summary")
    ac = [u for u in alloc.users if u.all_company]
    _table(ws, ["Username", "Emp Name", "Usage (GB)"], [[u.raw, u.person.emp_name, u.total] for u in ac], {2: GB},
           ["Total pooled usage", "", alloc.pool])
    ws.append([])
    ws.append(["Redistributed equally to users", alloc.recipients])
    ws.append(["Share per user (GB)", alloc.pool / alloc.recipients])

    ws = wb.create_sheet("Company Wise Summary")
    grp: "OrderedDict[str, list]" = OrderedDict()
    for u in alloc.users:
        if u.all_company:
            continue
        g = grp.setdefault(u.person.company or "(blank)", [0, Decimal(0), Decimal(0)])
        g[0] += 1
        g[1] += u.final
        g[2] += u.cost
    _table(ws, ["Company", "Users", "Final Usage (GB)", "Allocated Cost"],
           [[c, *v] for c, v in grp.items()], {2: GB, 3: MONEY},
           ["TOTAL", sum(v[0] for v in grp.values()), alloc.grand_total, alloc.allocated])

    period = re.sub(r"[^0-9A-Za-z]+", "-", bill.bill_period or "period").strip("-")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{bill.bill_type}_{bill.reference.replace(' ', '')}_{period}_Usage_Allocation.xlsx"
