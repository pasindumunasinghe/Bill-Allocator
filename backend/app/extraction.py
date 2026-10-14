"""Raw text extraction from uploaded documents.

Each reader returns plain text. No interpretation happens here; parsing lives in parser.py.
Reading is retried once; persistent failure raises ReadFailure with a human-readable reason.
"""
from __future__ import annotations

import os
import shutil
from dataclasses import dataclass, field

MIN_PAGE_CHARS = 20  # below this a PDF page is treated as scanned / image-only


class ReadFailure(Exception):
    pass


class UnsupportedFormat(Exception):
    pass


@dataclass
class ExtractionResult:
    text: str
    method: str
    warnings: list[str] = field(default_factory=list)
    attempts: int = 1


def ocr_available() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return shutil.which("tesseract") is not None


def _read_pdf(path: str) -> ExtractionResult:
    import pdfplumber

    texts, warnings, methods = [], [], set()
    with pdfplumber.open(path) as pdf:
        if not pdf.pages:
            raise ReadFailure("PDF contains no pages")
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if len(text.strip()) >= MIN_PAGE_CHARS:
                texts.append(text)
                methods.add("pdf-text")
                continue
            # Scanned page: OCR fallback
            if not ocr_available():
                warnings.append(f"Page {i} has no text layer (scanned) and OCR is not installed")
                continue
            import pytesseract

            image = page.to_image(resolution=300).original
            ocr_text = pytesseract.image_to_string(image)
            if ocr_text.strip():
                texts.append(ocr_text)
                methods.add("ocr")
                warnings.append(f"Page {i} read via OCR; verify values")
            else:
                warnings.append(f"Page {i} OCR produced no text")
    text = "\n".join(texts)
    if not text.strip():
        raise ReadFailure("; ".join(warnings) or "No readable text in PDF")
    return ExtractionResult(text, "+".join(sorted(methods)), warnings)


def _read_docx(path: str) -> ExtractionResult:
    import docx

    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip()
                if t and (not cells or cells[-1] != t):  # merged cells repeat text
                    cells.append(t)
            lines.append("  ".join(cells))
    text = "\n".join(lines)
    if not text.strip():
        raise ReadFailure("DOCX contains no text")
    return ExtractionResult(text, "docx")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_xlsx(path: str) -> ExtractionResult:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row if _cell(v)]
            if cells:
                lines.append("  ".join(cells))
    wb.close()
    text = "\n".join(lines)
    if not text.strip():
        raise ReadFailure("Workbook contains no data")
    return ExtractionResult(text, "xlsx")


def _read_xls(path: str) -> ExtractionResult:
    import xlrd

    book = xlrd.open_workbook(path)
    lines = []
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            cells = [_cell(v) for v in sheet.row_values(r) if _cell(v)]
            if cells:
                lines.append("  ".join(cells))
    text = "\n".join(lines)
    if not text.strip():
        raise ReadFailure("Workbook contains no data")
    return ExtractionResult(text, "xls")


READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".xlsx": _read_xlsx, ".xls": _read_xls}


def extract_text(path: str) -> ExtractionResult:
    ext = os.path.splitext(path)[1].lower()
    reader = READERS.get(ext)
    if reader is None:
        raise UnsupportedFormat(f"'{ext or 'no extension'}' files are not supported (use PDF, DOCX, XLSX or XLS)")
    last_error: Exception | None = None
    for attempt in (1, 2):  # retry once
        try:
            result = reader(path)
            result.attempts = attempt
            return result
        except Exception as exc:  # noqa: BLE001 - any parser error counts as a read failure
            last_error = exc
    raise ReadFailure(f"Unreadable after 2 attempts: {last_error}")
