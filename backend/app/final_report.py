"""Cumulative 'Final Report WiFi' workbook: one sheet per contract with a block per bill month,
plus a Company Summary (companies x months). Built from every successfully processed bill
across all batches; a later upload of the same contract + month replaces the earlier one."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import statuses as S
from .config import GROUP_A, GROUP_A_CONTRACT, GROUP_B_MAPPING
from .parser import parse_period

HEADER_FILL = PatternFill("solid", fgColor="1F3A68")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = "#,##0.00"
CONTRACT_ORDER = ["72263692", "72272015", GROUP_A_CONTRACT, "74746314", "B603042", "DBN0103018",
                  "DBN0160965", "DBN0235456", "ES00008395", "ES00014124"]
# Mapping names -> the company names used in the Company Summary
SUMMARY_NAME = {"Admin company": "OCP Admin Company", "Innovation Company": "OCP Innovation Company"}
SUMMARY_COMPANIES = ["MD's Office", "OCP Admin Company", "OCP Finance Company", "OCP Innovation Company",
                     "OCP Manufacturing Company", "OCP Product Company", "OCP Supply Chain Company"]
A_HEAD = ["Source File Name", "Contract Number", "Subscription Number", "Bill Period", "Bill Month",
          "Company", "Dimension", "Allocated Amount"]
B_HEAD = ["Source File Name", "Contract Number", "Subscription Number", "Bill Period", "Bill Month",
          "Related Company", "Cost Center", "Total Charges for Bill Period", "VAT", "Total Charges without VAT", "Status"]


def bill_month_key(f) -> date | None:
    if not f.bill_period or " - " not in f.bill_period:
        return None
    _, end = parse_period(*f.bill_period.split(" - ", 1))
    return date(end.year, end.month, 1) if end else None


def _months(first: date, last: date):
    y, m = first.year, first.month
    while (y, m) <= (last.year, last.month):
        yield date(y, m, 1)
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def _header(ws, headers):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT


def collect(files):
    """{contract: {month: file}} - latest processed file wins per contract+month."""
    out: dict[str, dict[date, object]] = {}
    for f in sorted((f for f in files if f.status == S.PROCESSED), key=lambda f: f.processed_at or f.created_at):
        key = bill_month_key(f)
        if key:
            out.setdefault(f.contract_number, {})[key] = f
    return out


def build_final_report(files) -> tuple[bytes, str]:
    data = collect(files)
    wb = Workbook()
    wb.remove(wb.active)
    all_months = sorted({m for per in data.values() for m in per})
    totals: dict[str, dict[date, Decimal]] = {}

    def add(company, month, amount):
        company = SUMMARY_NAME.get(company, company)
        totals.setdefault(company, {}).setdefault(month, Decimal(0))
        totals[company][month] += amount

    for contract in CONTRACT_ORDER:
        ws = wb.create_sheet(contract)
        per = data.get(contract, {})
        if not per:
            ws.append(["No processed bills yet"])
            continue
        for month in _months(min(per), max(per)):
            label = month.strftime("%B %Y")
            f = per.get(month)
            if f is None:
                ws.append([f"No bill value for {month.strftime('%B')}"])
                ws.append([])
                continue
            if contract == GROUP_A_CONTRACT:
                itemised = bool(f.allocations[0].get("subscription"))
                _header(ws, A_HEAD + (["Tax", "Full Amount"] if itemised else []))
                for a in f.allocations:
                    amt = Decimal(a["amount"])
                    if itemised:  # Allocated Amount = services, Tax = levies, Full Amount = excl. VAT
                        ws.append([f.file_name, f.contract_number, a["subscription"], f.bill_period, label,
                                   a["company"], a["dimension"], Decimal(a["charges"]), Decimal(a["levies"]), amt])
                    else:
                        ws.append([f.file_name, f.contract_number, f.subscription_number, f.bill_period, label,
                                   a["company"], a["dimension"], amt])
                    add(a["company"], month, amt)
            else:
                _header(ws, B_HEAD)
                ws.append([f.file_name, f.contract_number, f.subscription_number, f.bill_period, label,
                           f.related_company, f.cost_center, Decimal(str(f.total_charges)), Decimal(str(f.vat)),
                           Decimal(str(f.total_without_vat)), f.status])
                add(f.related_company, month, Decimal(str(f.total_without_vat)))
            ws.append([])
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, Decimal):
                    c.number_format = MONEY
        for i, w in enumerate([28, 16, 20, 25, 16, 28, 14, 28, 12, 26, 24], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet("Company Summary")
    _header(ws, ["Company"] + [m.strftime("%B %Y") for m in all_months] + ["Total"])
    companies = SUMMARY_COMPANIES + sorted(set(totals) - set(SUMMARY_COMPANIES))
    for comp in companies:
        vals = [totals.get(comp, {}).get(m, Decimal(0)) for m in all_months]
        ws.append([comp] + vals + [sum(vals, Decimal(0))])
    col_tot = [sum((totals.get(c, {}).get(m, Decimal(0)) for c in companies), Decimal(0)) for m in all_months]
    ws.append(["Grand Total"] + col_tot + [sum(col_tot, Decimal(0))])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row[1:]:
            c.number_format = MONEY
    ws.column_dimensions["A"].width = 30
    for i in range(2, len(all_months) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.append([])
    ws.append(["WiFi / fixed-line contracts only (amounts exclude VAT). Group B contracts with no monthly "
               "bill are shown as 'No bill value for <month>' on their sheet."])

    tag = all_months[-1].strftime("%b").upper() if all_months else "EMPTY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"Final_Report_WiFi_{tag}.xlsx"


assert set(CONTRACT_ORDER) == set(GROUP_B_MAPPING) | {GROUP_A_CONTRACT} and GROUP_A
