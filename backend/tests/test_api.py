import io
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from tests.make_samples import make


def test_end_to_end_mixed_batch(tmp_path):
    samples = make(tmp_path / "samples")
    with TestClient(app) as c:
        batch = c.post("/api/batches").json()
        files = [("files", (p.name, p.read_bytes())) for p in sorted(samples.iterdir())]
        batch = c.post(f"/api/batches/{batch['id']}/files", files=files).json()
        assert batch["summary"]["total_files"] == 13
        # export blocked until processed
        assert c.get(f"/api/batches/{batch['id']}/reports/group-a").status_code == 409
        for f in batch["files"]:
            assert c.post(f"/api/files/{f['id']}/process").status_code == 200
        batch = c.get(f"/api/batches/{batch['id']}").json()
        by = {f["file_name"]: f for f in batch["files"]}
        assert by["GroupB_ES00008395_corrupt.pdf"]["status"] == "File Read Failure"
        assert by["GroupB_ES00008395_corrupt.pdf"]["group"] == "B"  # routed by filename, not dropped
        assert by["GroupA_74745984_scanned.pdf"]["status"] == "File Read Failure"
        assert by["Invalid_contract_99999999.pdf"]["status"].startswith("Ignored")
        assert by["notes.txt"]["status"] == "Unsupported Format"
        s = batch["summary"]
        assert s["processed"] == 7 and s["group_a_processed"] == 2 and s["group_b_processed"] == 5

        ra = c.get(f"/api/batches/{batch['id']}/reports/group-a")
        assert "74745984_Allocation.xlsx" in ra.headers["content-disposition"]
        wa = openpyxl.load_workbook(io.BytesIO(ra.content))
        assert wa.sheetnames == ["Allocation Summary", "Calculation Summary", "Validation"]
        assert wa["Allocation Summary"].max_row == 1 + 8
        assert all(r[1] == "74745984" for r in wa["Allocation Summary"].iter_rows(min_row=2, values_only=True))
        statuses = {r[0]: r[2] for r in wa["Validation"].iter_rows(min_row=2, values_only=True)}
        assert statuses["GroupA_74745984_scanned.pdf"] == "File Read Failure"

        rb = c.get(f"/api/batches/{batch['id']}/reports/group-b")
        assert "Telecom_Bill_Summary.xlsx" in rb.headers["content-disposition"]
        wb = openpyxl.load_workbook(io.BytesIO(rb.content))
        assert wb.sheetnames == ["Bill Summary", "Cost Center Summary", "Failed Records"]
        contracts = [r[1] for r in wb["Bill Summary"].iter_rows(min_row=2, values_only=True)]
        assert "74745984" not in contracts and len(contracts) == 5
        cc = {r[0]: r[2] for r in wb["Cost Center Summary"].iter_rows(min_row=2, values_only=True)}
        assert float(cc["1100700610"]) == 800.00  # 500 + 300
        failed = [r[0] for r in wb["Failed Records"].iter_rows(min_row=2, values_only=True)]
        assert set(failed) == {"GroupB_74746314_missing_VAT.pdf", "GroupB_DBN0235456_payable_only.pdf",
                               "GroupB_ES00008395_corrupt.pdf"}
