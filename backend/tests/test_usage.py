import io
from decimal import Decimal

import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.usage.usage import norm, raw_user
from tests.make_usage_samples import make


def test_username_normalization():
    assert raw_user("OREL\\Kamal ") == "kamal"
    assert norm("Nimal.P@orel.com") == "nimalp" and norm("sunil_k") == norm("sunil-k") == "sunilk"


def _run(c, s, bills, **extra):
    files = [("bills", (b, (s / b).read_bytes())) for b in bills]
    files += [("usage", ("usage.xlsx", (s / "usage.xlsx").read_bytes())),
              ("master", ("Emp.xlsx", (s / "Emp.xlsx").read_bytes())),
              ("map_file", ("map.xlsx", (s / "map.xlsx").read_bytes()))]
    return c.post("/api/usage/allocate", files=files, **extra)


def test_two_bills_two_files(tmp_path):
    s = make(tmp_path)
    with TestClient(app) as c:
        r = _run(c, s, ["Dialog_B698359_Jun.pdf", "SLT_0033118626_Jun.pdf"])
        assert r.status_code == 200, r.text
        res = {x["bill_type"]: x for x in r.json()["results"]}
        assert res["DIALOG"]["bill_value"] == "318228.86"
        assert res["SLT"]["bill_value"] == "140704.92"
        for x in res.values():
            assert x["status"] == "Processed Successfully" and x["allocated"] == x["bill_value"]
        wb = openpyxl.load_workbook(io.BytesIO(c.get(res["SLT"]["download"]).content))
        assert wb.sheetnames == ["Allocation Summary", "Unmapped Users", "Bill Summary",
                                 "All Company Summary", "Company Wise Summary"]
        rows = {r[0]: r for r in wb["Allocation Summary"].iter_rows(min_row=2, values_only=True) if r[0]}
        # kamal: 3 GB (sheet1 MB) + 1 GB; nimal: 1 + 2 (first occurrence only); sunil via map: 4 + 1; shared 0.5 pooled
        assert Decimal(str(rows["kamal"][10])) == Decimal("4")
        assert Decimal(str(rows["nimal.p"][10])) == Decimal("3")
        assert rows["sunil_k"][6].startswith("map.xlsx")
        unmapped = [r[0] for r in wb["Unmapped Users"].iter_rows(min_row=2, values_only=True)]
        assert unmapped == ["ghost"]
        assert abs(sum(Decimal(str(r[3])) for r in wb["Company Wise Summary"].iter_rows(min_row=2, max_row=4, values_only=True))
                   - Decimal("140704.92")) < Decimal("0.001")


def test_more_than_two_bills_rejected(tmp_path):
    s = make(tmp_path)
    with TestClient(app) as c:
        r = _run(c, s, ["Dialog_B698359_Jun.pdf", "SLT_0033118626_Jun.pdf", "Other_bill.pdf"])
        assert r.status_code == 400 and "maximum 2" in r.text


def test_invalid_bill_skipped(tmp_path):
    s = make(tmp_path)
    with TestClient(app) as c:
        res = _run(c, s, ["Other_bill.pdf"]).json()["results"][0]
        assert res["status"] == "Invalid" and "download" not in res


def test_usage_bill_on_upload_page_is_redirected(tmp_path):
    s = make(tmp_path)
    with TestClient(app) as c:
        b = c.post("/api/batches").json()
        files = [("files", (n, (s / n).read_bytes())) for n in ["SLT_0033118626_Jun.pdf", "Dialog_B698359_Jun.pdf"]]
        b = c.post(f"/api/batches/{b['id']}/files", files=files).json()
        b = c.post(f"/api/batches/{b['id']}/process").json()
        for f in b["files"]:
            assert f["status"] == "Usage Allocation Bill" and not f["errors"]
        assert b["summary"]["failed"] == 0


def test_preview_bill_and_allocation_preview(tmp_path):
    s = make(tmp_path)
    with TestClient(app) as c:
        p = c.post("/api/usage/preview-bill", files={"bill": ("d.pdf", (s / "Dialog_B698359_Jun.pdf").read_bytes())}).json()
        assert p["bill_type"] == "DIALOG" and p["bill_value"] == "318228.86" and len(p["components"]) == 2
        r = _run(c, s, ["SLT_0033118626_Jun.pdf"]).json()["results"][0]["preview"]
        assert r["summary"]["allocated"] == "140704.92" and r["unmapped"][0][0] == "ghost"


def test_latest_by_bill(tmp_path):
    s = make(tmp_path)
    with TestClient(app) as c:
        _run(c, s, ["SLT_0033118626_Jun.pdf", "Dialog_B698359_Jun.pdf"])
        latest = c.get("/api/usage/latest").json()
        assert latest["SLT"]["bill_value"] == "140704.92" and latest["DIALOG"]["preview"]["summary"]["allocated"] == "318228.86"


def test_inputs_are_stored_and_previewed(tmp_path):
    s = make(tmp_path)
    up = lambda c, kind, f: c.post(f"/api/usage/inputs/{kind}", files={"file": (f, (s / f).read_bytes())})  # noqa: E731
    with TestClient(app) as c:
        assert up(c, "slt", "Dialog_B698359_Jun.pdf").status_code == 422  # wrong box, not stored
        m = up(c, "slt", "SLT_0033118626_Jun.pdf").json()
        assert m["stored_at"] and m["preview"]["bill_value"] == "140704.92"
        u = up(c, "usage", "usage.xlsx").json()["preview"]
        assert len(u["sheets"]) == 3 and u["users"] >= 4
        assert up(c, "master", "Emp.xlsx").json()["preview"]["employees"] == 3
        up(c, "map", "map.xlsx")
        cur = c.get("/api/usage/inputs").json()
        assert cur["slt"]["file_name"] == "SLT_0033118626_Jun.pdf" and cur["dialog"] is None
        r = c.post("/api/usage/allocate-stored").json()["results"]
        assert len(r) == 1 and r[0]["allocated"] == "140704.92"
        c.delete("/api/usage/inputs/slt")
        assert c.get("/api/usage/inputs").json()["slt"] is None
